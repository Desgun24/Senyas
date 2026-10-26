"""Create the requested preliminary report and feature derivation in plain Word formatting."""
from pathlib import Path
import json

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs/preliminary-validation-20261026'
CREDITS = 'Jeff Barrientos, Hannah Dimaculangan, and Trixie Mendoza'


def document(title, opening):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.72)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number']:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name.startswith('Heading')
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.08
    doc.styles['Title'].font.size = Pt(18)
    doc.styles['Heading 1'].font.size = Pt(13)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
    doc.styles['Heading 2'].font.size = Pt(11)
    doc.core_properties.author = CREDITS
    doc.core_properties.title = title
    doc.core_properties.subject = 'Senyas Filipino Sign Language research documentation'
    doc.add_paragraph(title, 'Title')
    doc.add_paragraph(CREDITS)
    doc.add_paragraph('26 October 2026')
    doc.add_paragraph(opening)
    return doc


def heading(doc, title):
    doc.add_heading(title, level=1)


def p(doc, text):
    doc.add_paragraph(text)


def eq(doc, text):
    paragraph = doc.add_paragraph()
    math = OxmlElement('m:oMath')
    run = OxmlElement('m:r')
    prop = OxmlElement('m:rPr')
    normal = OxmlElement('m:nor')
    prop.append(normal)
    run.append(prop)
    node = OxmlElement('m:t')
    node.text = text
    run.append(node)
    math.append(run)
    paragraph._p.append(math)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    if widths:
        for col, width in zip(t.columns, widths):
            col.width = Inches(width)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = str(text)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = str(text)
    for row in t.rows:
        for i, cell in enumerate(row.cells):
            if widths:
                cell.width = Inches(widths[i])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.space_before = Pt(4)
                for run in paragraph.runs:
                    run.font.name = 'Arial'
                    run.font.size = Pt(10.5)
                    run.font.color.rgb = RGBColor(0, 0, 0)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                border = OxmlElement(f'w:{edge}')
                for key, value in [('val', 'single'), ('sz', '4'), ('color', '000000')]:
                    border.set(qn('w:' + key), value)
                borders.append(border)
            cell._tc.get_or_add_tcPr().append(borders)
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def save(doc, name):
    # Remove default theme colours, fills, highlights, and decorative borders.
    for root in [doc._element, doc.styles.element]:
        for tag in ['w:shd', 'w:highlight', 'w:background', 'w:pBdr']:
            for node in list(root.iter(qn(tag))):
                node.getparent().remove(node)
        for color in root.iter(qn('w:color')):
            color.attrib.clear()
            color.set(qn('w:val'), '000000')
    target = OUT / name
    doc.save(target)
    print(target)


def feature_document(meta, report=None, example=None):
    doc = document('Senyas Engineered Feature Derivation',
        'This document explains how Senyas converts sampled signing video frames into 1,630 numerical inputs and how the Decision Tree uses those inputs. The formulas describe the current implementation in senyas/features.py and senyas/model.py. They explain measured geometry and learned thresholds; they do not define a complete linguistic grammar of Filipino Sign Language.')
    heading(doc, 'Representation and processing')
    p(doc, 'A clip is sampled at up to 32 approximately evenly spaced frame positions. Frames are resized only when their longest side exceeds 640 pixels, converted from BGR to RGB, and passed to MediaPipe Holistic. The configured detector uses static image mode, model complexity one, and detection confidence 0.5. Static image mode treats sampled frames independently; enabling smooth_landmarks does not establish continuous temporal tracking in that mode. Live camera attempts use the same descriptor functions after browser capture.')
    p(doc, 'Each frame has 163 base features. Every base feature receives ten temporal summaries in a fixed order, giving 163 × 10 = 1,630 inputs. Four base features are presence indicators. Frame rate, clip duration, extraction time, and detection rate are diagnostic metadata and do not enter the classifier vector.')
    table(doc, ['Feature family', 'Per frame', 'After summaries'], [
        ('Pose, left hand, right hand, and both-hands presence', 4, 40),
        ('Shoulder width', 1, 10), ('Absolute wrist coordinates for both hands', 4, 40),
        ('Body-relative wrist coordinates for both hands', 4, 40), ('Palm scales for both hands', 2, 20),
        ('Hand-relative landmark coordinates: 2 × 21 × 3', 126, 1260),
        ('Finger bend measures: 2 × 5', 10, 100), ('Fingertip distances: 2 × 5', 10, 100),
        ('Distance and vertical offset between wrists', 2, 20), ('Total', 163, 1630)], [4.9, .9, 1.2])
    p(doc, 'MediaPipe is a learned landmark detector upstream of the transparent classifier. The Decision Tree decision is inspectable given its inputs, but the complete pixels-to-word pipeline is not entirely white box. Face landmarks are not converted into classifier features, and the code does not calculate a palm normal or an explicit inward/outward-facing label.')

    heading(doc, 'Coordinate conventions and shoulder reference')
    p(doc, 'For hand h, let P(h,j,t) = (x,y,z) be landmark j at sampled frame t. Image x increases to the right and image y increases downward. MediaPipe x and y are normalized by image dimensions, while hand z is an estimated relative depth quantity. These are not calibrated physical distances in metres. The current 2D norms mix normalized x and y directly, so camera aspect ratio and viewpoint can affect them. Left and right refer to detector-assigned anatomical hands; the mirrored preview does not mirror the analyzed frame.')
    p(doc, 'Shoulder landmarks are pose indices 11 and 12. Their 2D midpoint is the body reference C(t); their gap is S(t). The lower bound prevents division by a tiny number.')
    eq(doc, 'C(t) = (P(pose,11,t)[xy] + P(pose,12,t)[xy]) / 2')
    eq(doc, 'S(t) = max( ‖P(pose,11,t)[xy] − P(pose,12,t)[xy]‖₂ , 10⁻⁴ )')
    p(doc, 'pose_present is one when pose landmarks exist and zero otherwise. shoulder_width stores S(t). If the pose is absent, C(t) becomes (0.5, 0.5) and S(t) becomes 1.0. A wrist_body value then refers to the fallback image centre, so it cannot truthfully be explained as a measured shoulder-relative position. Pose visibility scores are not used to reject individual shoulders in this implementation.')
    heading(doc, 'Absolute and body relative wrist position')
    p(doc, 'For each hand, present is one when its landmarks exist. wrist_abs_x and wrist_abs_y store the original image-normalized wrist coordinates. wrist_body_x and wrist_body_y express the wrist displacement from the shoulder midpoint in units of the visible shoulder gap.')
    eq(doc, 'W(h,t) = P(h,0,t);  B(h,t) = (W(h,t)[xy] − C(t)) / S(t)')
    p(doc, 'Positive wrist_body_y means below the reference; negative means above it. Positive wrist_body_x means image-right of the reference. These features capture hand location near the torso; they are not finger-shape features and do not establish body pose or facial expression.')

    heading(doc, 'Hand scale and the landmark coordinates')
    p(doc, 'The palm_scale feature is the 2D distance from the wrist to the middle-finger base knuckle, bounded below by 0.0001. It is a visible image-space scale proxy, not a physical palm measurement. The normalized vector Q subtracts the wrist and divides all three coordinates by this scale.')
    eq(doc, 's(h,t) = max( ‖P(h,9,t)[xy] − P(h,0,t)[xy]‖₂ , 10⁻⁴ )')
    eq(doc, 'Q(h,j,t) = (P(h,j,t) − P(h,0,t)) / s(h,t)')
    p(doc, 'Each left or right landmark contributes _x, _y, and _z. These coordinates retain hand shape and orientation relative to the wrist while reducing translation and apparent size effects. No rotation alignment is applied, so viewpoint and hand rotation remain relevant. The normalized wrist itself is always zero when observed; its six coordinate streams remain zero after missing-value handling and therefore provide no discriminating signal.')
    table(doc, ['Indices', 'Landmark names in feature keys'], [
        ('0', 'wrist'), ('1, 2, 3, 4', 'thumb_cmc, thumb_mcp, thumb_ip, thumb_tip'),
        ('5, 6, 7, 8', 'index_mcp, index_pip, index_dip, index_tip'),
        ('9, 10, 11, 12', 'middle_mcp, middle_pip, middle_dip, middle_tip'),
        ('13, 14, 15, 16', 'ring_mcp, ring_pip, ring_dip, ring_tip'),
        ('17, 18, 19, 20', 'pinky_mcp, pinky_pip, pinky_dip, pinky_tip')], [1.3, 5.7])
    p(doc, 'CMC denotes the thumb base, MCP the base knuckle, PIP a middle finger joint, DIP an end finger joint, and IP the thumb interphalangeal joint. For example, right.index_tip_y.mean describes the average vertical position of the right index fingertip relative to its wrist and palm scale, not the average height of the whole hand.')

    heading(doc, 'Finger bend and fingertip extension')
    p(doc, 'For a selected triple of normalized landmarks (a,b,c), the implementation measures the angle at b and divides by π. It clips the cosine into [−1,1] before arccos to avoid numerical error. A value near one is straighter at the measured joint; smaller values indicate a sharper bend. These are continuous features, not preassigned bent/straight categories.')
    eq(doc, 'u = Q(a) − Q(b);  v = Q(c) − Q(b)')
    eq(doc, 'bend = arccos( clip( (u · v) / (‖u‖₂ ‖v‖₂), −1, 1 ) ) / π')
    p(doc, 'If the product of the vector norms is below 10⁻⁸, the implementation returns zero. That is a numerical fallback and cannot by itself prove a bent finger. Four finger angles use a base-to-middle-to-tip configuration; the thumb uses its MCP, IP, and tip. This does not measure every anatomical joint separately.')
    table(doc, ['Finger', 'Angle indices a b c', 'Tip index'], [('Thumb', '2, 3, 4', 4), ('Index', '5, 6, 8', 8), ('Middle', '9, 10, 12', 12), ('Ring', '13, 14, 16', 16), ('Pinky', '17, 18, 20', 20)], [2, 3.5, 1.5])
    p(doc, 'The tip_distance feature is the 2D wrist-to-fingertip distance in normalized hand coordinates. A larger distance usually describes greater visible extension, but rotation and foreshortening can also change it. The angle uses three coordinates; fingertip distance deliberately uses x and y only.')
    eq(doc, 'tip_distance(h,f,t) = √( Q(h,tip(f),t)[x]² + Q(h,tip(f),t)[y]² )')

    heading(doc, 'Relations between the hands')
    p(doc, 'both_hands_present is one only when both hand landmark sets exist in the same frame. When both wrists exist, hands_distance measures their 2D separation divided by shoulder width, and hands_vertical_offset measures left-wrist y minus right-wrist y in the same units.')
    eq(doc, 'hands_distance(t) = ‖W(left,t)[xy] − W(right,t)[xy]‖₂ / S(t)')
    eq(doc, 'hands_vertical_offset(t) = (W(left,t)[y] − W(right,t)[y]) / S(t)')
    p(doc, 'A positive vertical offset places the left wrist lower in the image than the right wrist; a negative offset places it higher. Both measurements are missing when either hand is absent. They use the fallback reference if the shoulders are unavailable, which must be disclosed when generating a human explanation.')

    heading(doc, 'Missing detections and sample acceptance')
    p(doc, 'A missing hand contributes present = 0 and missing values for its other 78 fields. A missing pair contributes both_hands_present = 0 and missing inter-hand distances. Presence flags remain observed zeros and ones; they are not interpolated into pose evidence. Clips must contain a detected left or right hand in at least three sampled frames to be accepted.')
    p(doc, 'For every scalar coordinate or geometry series, finite observations are interpolated over sampled-frame indices. Internal gaps receive linear interpolation. Missing values before the first observation use its value, and trailing gaps use the last observation. If an entire column has no finite observation, every value becomes zero. Frames are retained rather than discarded; a one-hand frame remains usable.')
    eq(doc, 'x̂(t) = x(a) + (t − a)(x(b) − x(a)) / (b − a), for a < t < b')
    p(doc, 'This equation describes internal gaps with observed endpoints a and b. Interpolation does not verify what the hand actually did during the gap. A zero-filled absent hand can otherwise look like a centred or tightly folded hand, so developer explanations use presence summaries to suppress unsupported anatomical claims. Detection-rate metadata reports frames containing either hand divided by sampled frames.')

    heading(doc, 'The ten temporal summaries')
    p(doc, 'For each filled scalar sequence x₁,…,xT, the following statistics are appended in this exact order. Standard deviation uses the population convention. Adjacent-change statistics use absolute differences and are not divided by elapsed time.')
    table(doc, ['Suffix', 'Formula or definition'], [
        ('mean', 'μ = (Σ xₜ) / T'), ('std', '√[Σ(xₜ − μ)² / T]'), ('min', 'minₜ xₜ'), ('max', 'maxₜ xₜ'),
        ('range', 'maxₜ xₜ − minₜ xₜ'), ('start', 'x₁'), ('end', 'xT'), ('delta', 'xT − x₁'),
        ('velocity_mean', 'Σ |xₜ₊₁ − xₜ| / (T − 1)'), ('velocity_max', 'maxₜ |xₜ₊₁ − xₜ|')], [1.7, 5.3])
    p(doc, 'For a one-frame sequence, both adjacent-change summaries are zero. For ordinary clips, velocity_mean is average change per sampled step, not metres per second or frames per second. A zero delta means matching endpoints and can coexist with considerable intermediate movement. A small standard deviation applies only to that scalar feature, not the steadiness of the entire hand.')
    p(doc, 'For binary presence streams, mean is detection coverage, min means presence in every sampled frame, max means presence at least once, and start/end report the endpoint detections. Their adjacent-change terms record switching between detected and undetected states. Missing-data handling and varying clip durations limit comparisons between these summaries.')
    p(doc, 'Feature ordering is base-major: all ten summaries for pose_present come first, followed by all ten for shoulder_width, then left-hand fields, right-hand fields, and inter-hand fields. If a base field has zero-based index b and its statistic has index s, the final vector index is 10b + s. The companion feature_catalog.csv lists every exact key and index.')

    heading(doc, 'How the Decision Tree learns a split')
    p(doc, 'Each internal node chooses one feature j and a threshold θ. A sample follows the left branch when x[j] ≤ θ and the right branch when x[j] > θ. Split thresholds are learned from training samples; the code does not contain a handcrafted rule such as a bent index finger always means one particular word. At a leaf, the predicted word is the class with the largest weighted class share.')
    eq(doc, 'Left: x[j] ≤ θ;  Right: x[j] > θ;  predicted class = argmaxₖ p(leaf,k)')
    p(doc, 'Training uses balanced class weights. If class k has nₖ training samples among N samples and K classes, each sample of that class receives weight N/(K nₖ). At a node, class proportions are computed from those weights. The candidate search uses either Gini impurity or entropy and selects splits that reduce weighted impurity.')
    eq(doc, 'wₖ = N / (K nₖ);  Gini = 1 − Σₖ pₖ²;  Entropy = −Σₖ pₖ log₂(pₖ)')
    eq(doc, 'Gain = I(parent) − (Wleft/Wparent)I(left) − (Wright/Wparent)I(right)')
    p(doc, 'Maximum depth restricts how many branching decisions are allowed. min_samples_leaf requires a minimum number of samples in a terminal region. Cost-complexity pruning trades impurity against the number of leaves through Rα(T) = R(T) + α|leaves(T)|. Thus a larger ccp_alpha can remove branches that give little benefit. The stored leaf vote fraction is not a calibrated probability that a learner signed correctly.')
    p(doc, 'The preliminary selection checks 36 combinations: two criteria, depths 6/10/14, minimum leaf sizes 2/4/8, and pruning values 0/0.001 (2 × 3 × 3 × 2 = 36). Every candidate is fitted on the training split. Validation macro F1 ranks candidates, then validation accuracy, then fewer nodes. Validation therefore informs parameter selection and the resulting scores are preliminary selection scores.')
    p(doc, f'The app currently contains a separate final-refit Decision Tree with {meta["app_tree_nodes"]} nodes, {meta["app_tree_leaves"]} leaves, depth {meta["app_tree_depth"]}, and {meta["app_tree_features_used"]} distinct split features. Its parameters are {meta["app_tree_parameters"]}. It was refitted after validation selection and is distinct from the train-only preliminary model.')
    if report and example:
        heading(doc, 'Verified decision path from this preliminary run')
        p(doc, f'The example is {example["vid_path"]}; its true class is {example["true_id"]} and its predicted word is {example["prediction"]["label"]}. The following rows are successive comparisons on one complete path. Each comparison was checked against the saved child node; the leaf prediction agrees with the model output. A correct path trace documents what the model did, even when a classification is wrong.')
        rows = [(r['feature'], f'{r["value"]:.5f}', f'{r["operator"]} {r["threshold"]:.5f}') for r in example['rules']]
        table(doc, ['Engineered feature', 'Observed value', 'Chosen branch'], rows, [4.25, 1.2, 1.55])
        p(doc, 'In this example, the zero hands_distance.max comes from filling a sequence with no simultaneous two-hand detection. The first branch is therefore a real numerical decision, but it must not be described as proof that the hands were close together.')
    heading(doc, 'Human explanations and known limits')
    p(doc, 'Developer mode translates comparisons into qualitative descriptions tied to the relevant coordinate reference and temporal statistic. It retains the exact numeric comparisons internally. A phrase such as a more bent index finger is a description relative to a learned threshold, not a clinical joint measurement or proof of a linguistic cause. Missing or partial tracking must qualify the description, and old logs without coverage metadata require a fresh recording for verified hand-state descriptions.')
    p(doc, 'With Decision Tree selected, the displayed path belongs to the recognizer itself. With Extra Trees selected, the separate Decision Tree baseline cannot explain every ensemble vote and is labeled accordingly. SHAP is not currently part of this implementation. Face expression, gaze, explicit palm normals, full-body posture, physical speed, and sentence-level grammar are outside the current engineered feature set.')
    p(doc, 'Translation and scale normalization do not eliminate camera-angle effects, occlusion, detector errors, or image-aspect distortions. Temporal summary vectors discard much of the order of intermediate poses; different trajectories can produce similar summaries. Presence and scale features can also capture recording conditions rather than sign semantics. Interpretability of a path does not establish causal validity or robustness for every learner.')
    heading(doc, 'Implementation and references')
    for text in [
        'Implementation: senyas/features.py, senyas/model.py, senyas/explanations.py, and scripts/manifest_validation.py in the accompanying Senyas source. Feature metadata was read from the current artifact and schema on 26 October 2026.',
        'Feature registry: feature_catalog.csv contains all 1,630 exact names, indices, and final app-tree usage indicators. Impurity importance records split contribution in that fitted tree; it is not causal importance.',
        'MediaPipe Holistic documentation: https://chuoling.github.io/mediapipe/solutions/holistic.html',
        'MediaPipe Hands coordinate documentation: https://chuoling.github.io/mediapipe/solutions/hands.html',
        'scikit-learn 1.6 Decision Trees: https://scikit-learn.org/1.6/modules/tree.html',
        'scikit-learn 1.6 DecisionTreeClassifier: https://scikit-learn.org/1.6/modules/generated/sklearn.tree.DecisionTreeClassifier.html']:
        p(doc, text)
    save(doc, 'Senyas Feature Derivation.docx')


def validation_document(r, audit):
    m = r['validation_metrics_extracted']
    doc = document('Senyas Preliminary Validation Results',
        f'A locally trained Decision Tree correctly classified {r["correct"]} of {r["validation_extracted"]} validation clips that produced valid features, giving {m["accuracy"]:.2%} accuracy and {m["macro_f1"]:.4f} macro F1. Validation was used to select model parameters, so these are preliminary selection results. The separate test split was not processed or scored in this run.')
    heading(doc, 'Run scope and data assignment')
    p(doc, 'The supplied manifest.csv determines split membership using its split column. Its source_split column is historical metadata and was not used to reassign samples. Original relative clip paths were matched to the local D:\\DATASETS_FSL collection; the other computer’s absolute paths were discarded. Numeric class IDs and category assignments were preserved.')
    table(doc, ['Split', 'Clips', 'Classes', 'Signer or session IDs'], [(name.title(), part['samples'], part['classes'], ', '.join(map(str, part['signer_ids']))) for name, part in audit['splits'].items()], [1, .7, .9, 4.4])
    p(doc, 'All 2,130 partition copies were verified byte-for-byte using SHA-256. No signer/session ID occurs in more than one split, and no identical video bytes cross split boundaries. Filename-derived IDs are treated as grouping identifiers; the manifest alone does not prove that every ID represents a unique person. The only source-label discrepancy was whitespace after NINE, which did not change its numeric ID.')
    heading(doc, 'Feature extraction coverage')
    table(doc, ['Measure', 'Result'], [('Training clips accepted', r['train_samples']), ('Validation clips attempted', r['validation_attempted']), ('Validation clips accepted', r['validation_extracted']), ('Validation clips rejected', r['validation_failed']), ('Test clips processed', r['test_samples_processed']), ('Model features per accepted clip', r['feature_count'])], [4.7, 2.3])
    p(doc, 'The extractor samples up to 32 frames, uses MediaPipe Holistic hand and shoulder landmarks, and computes 163 frame features with ten temporal summaries. A clip must show at least one detected hand in three sampled frames. Missing scalar observations are interpolated within a clip; completely missing columns become zero. The same feature code is used by the app, although prerecorded clips and webcam capture can have different recording conditions.')
    import csv
    with (OUT / 'extraction_failures.csv').open(newline='', encoding='utf-8') as stream:
        failures = list(csv.DictReader(stream))
    for item in failures:
        p(doc, f'Rejected clip: {item["vid_path"]} ({item["split"]}, {item["label"]}). Reason: {item["error"]}. This clip remains in the manifest and all-attempted denominator.')
    heading(doc, 'Training and parameter selection')
    p(doc, f'A fresh DecisionTreeClassifier was fitted on training samples only. The existing search evaluated {r["candidate_count"]} combinations of criterion, maximum depth, minimum leaf size, and cost-complexity pruning. Candidates were ordered by validation macro F1, then validation accuracy, then fewer nodes. Balanced class weights and random seed 42 were used. The selected candidate was saved directly without refitting on validation data.')
    table(doc, ['Selected setting', 'Value'], [(k, v) for k, v in r['selected_parameters'].items()] + [('class_weight', 'balanced'), ('random_state', 42), ('actual tree depth', r['tree_depth']), ('leaves', r['tree_leaves']), ('distinct split features', r['used_features'])], [4.7, 2.3])
    p(doc, 'The app artifact was left unchanged. Its original training script performs candidate selection first, then refits the selected models on training plus validation before final test evaluation. Scoring that refitted artifact on validation would reuse fitting samples; this preliminary report instead evaluates the newly fitted training-only candidate.')
    heading(doc, 'Preliminary validation metrics')
    table(doc, ['Metric', 'Accepted validation clips'], [
        ('Accuracy', f'{m["accuracy"]:.4f} ({m["accuracy"]:.2%})'), ('Macro precision', f'{m["macro_precision"]:.4f}'),
        ('Macro recall', f'{m["macro_recall"]:.4f}'), ('Macro F1', f'{m["macro_f1"]:.4f}'),
        ('Weighted precision', f'{r["weighted_metrics"]["precision"]:.4f}'), ('Weighted recall', f'{r["weighted_metrics"]["recall"]:.4f}'),
        ('Weighted F1', f'{r["weighted_metrics"]["f1-score"]:.4f}'), ('Correct predictions', f'{r["correct"]} / {r["validation_extracted"]}')], [4.7, 2.3])
    p(doc, f'Training accuracy was {r["training_accuracy"]:.2%}. The gap from validation accuracy shows that fitting the training examples did not translate into equally strong performance on the validation groups.')
    p(doc, f'Counting extraction failures as unsuccessful attempts gives {r["correct"]}/{r["validation_attempted"]} = {r["accuracy_all_attempted"]:.2%} accuracy across the complete validation manifest. The two accuracy figures answer different questions: accepted-clip accuracy measures classification after feature extraction, while all-attempted accuracy includes extraction failures. No confidence rejection threshold was imposed on the Decision Tree.')
    p(doc, 'Macro metrics give each of the 105 word classes equal weight. Weighted metrics weight each class by its validation support. Precision is TP/(TP+FP), recall is TP/(TP+FN), and F1 is 2PR/(P+R); undefined class terms are assigned zero. Accuracy is the number of correct predictions divided by the stated denominator.')
    heading(doc, 'Validation groups')
    table(doc, ['Category', 'Attempted', 'Accepted', 'Correct', 'Accuracy'], [(a['group'], a['attempted'], a['extracted'], a['correct'], f'{a["accuracy_all_attempted"]:.2%}') for a in r['groups']['category']], [2.4, 1.1, 1.1, 1.1, 1.3])
    table(doc, ['Signer or session ID', 'Attempted', 'Accepted', 'Correct', 'Accuracy'], [(a['group'], a['attempted'], a['extracted'], a['correct'], f'{a["accuracy_all_attempted"]:.2%}') for a in r['groups']['signer_id']], [2.4, 1.1, 1.1, 1.1, 1.3])
    p(doc, 'Group accuracy includes extraction failures. Small groups, particularly the incomplete signer/session group, should not be interpreted as stable population estimates. No lighting, distance, background, or speed labels were available for a controlled condition comparison.')
    heading(doc, 'Decision path verification and interpretation')
    p(doc, f'All {r["trace_checks"]} accepted validation decisions were checked against the actual tree child sequence and leaf prediction, with {r["valid_traces"]} valid paths. This verifies trace consistency, not classification correctness. The accompanying feature-derivation document explains how each input relates to landmarks and why missing observations restrict human-readable claims.')
    p(doc, 'These scores are preliminary because the validation set influenced parameter selection. They are not a final held-out estimate and should not be presented as webcam accuracy or complete FSL understanding. A later final evaluation should follow a frozen protocol and use the separate test split once the model and reporting choices are fixed. The 315 test clips remain available in their partition folder without inference in this run.')
    heading(doc, 'Reproducibility and deliverables')
    p(doc, 'The local partition root is D:\\DATASETS_FSL\\partitions\\supplied-manifest-20261026. manifest_local.csv points to the copied clips; manifest_portable.csv preserves relative paths for another computer. scripts/manifest_validation.py provides prepare, extract, and evaluate commands. Feature caching is keyed to source video bytes, feature-code hash, extractor settings, and the pinned MediaPipe version.')
    p(doc, 'The evidence archive contains metrics.json, validation_predictions.csv, validation_per_class.csv, validation_confusion_matrix.csv, decision_tree_candidates.csv, extraction_failures.csv, split_audit.json, manifest_portable.csv, feature_catalog.csv, the verified example path, and the separate preliminary Decision Tree artifact. The failure row is retained in validation_predictions.csv.')
    table(doc, ['Provenance item', 'SHA 256'], [('Supplied manifest', r['source_manifest_sha256']), ('Feature implementation', r['features_source_sha256']), ('Preliminary model', r['preliminary_model_sha256']), ('Unchanged app model', r['app_model_sha256_after'])], [1.6, 5.4])
    p(doc, 'Environment: Python 3.11, MediaPipe 0.10.21, NumPy 1.26.4, and scikit-learn 1.6.1. Runtime measurements, exact predictions, and complete class-level counts are retained in the machine-readable evidence rather than inferred from rounded percentages.')
    save(doc, 'Senyas Preliminary Validation Results.docx')


if __name__ == '__main__':
    meta = json.loads((OUT / 'feature_metadata.json').read_text())
    report = json.loads((OUT / 'metrics.json').read_text()) if (OUT / 'metrics.json').exists() else None
    example = json.loads((OUT / 'example_decision_path.json').read_text()) if report else None
    feature_document(meta, report, example)
    if report:
        validation_document(report, json.loads((OUT / 'split_audit.json').read_text()))
